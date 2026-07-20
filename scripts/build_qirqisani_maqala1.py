"""
Build per-bab bundles for Qirqisānī Maqala I (babs 2–19) for the Advanced reader.

Inputs (outside the repo):
  - English + segmentation map:
      ~/Downloads/genizah-research/qirqisani-maqala1-english/qirqisani_maqala1_parallel.xlsx
      sheet "Text Content": Index col "I.{bab}.{seg}" (or "I.{bab}" for bab headers),
      English col (carries [{n}] footnote markers we strip).
  - Arabic witness (noisy OCR, clean reading order, mapped bab boundaries):
      ~/Downloads/الأنوار والمراقب Arabic ed .txt   (Cairo 2019)

Outputs (build artifacts, NOT shipped):
  - data/_qirqisani_build/bab{N}.json  for N in 2..19:
      { bab, title_en, n_segments, segments:[{seg, en}], cairo_arabic }
    Consumed by the reconstruction step, which produces clean Arabic per segment.
  - data/_qirqisani_footnotes.json     (Footnotes sheet dumped for later; sidecar)
  - data/_qirqisani_english_skeleton.json (all 19 babs, English only, for reference)

Bab 1 is NOT rebuilt — the existing hand-built opening in
data/qirqisani-anwar-maqala1.json (page 1) is preserved as gold.

Run from repo root: python3 scripts/build_qirqisani_maqala1.py
"""
import json
import pathlib
import re
import sys

HOME = pathlib.Path.home()
XLSX = HOME / "Downloads" / "genizah-research" / "qirqisani-maqala1-english" / "qirqisani_maqala1_parallel.xlsx"
CAIRO = HOME / "Downloads" / "الأنوار والمراقب Arabic ed .txt"

REPO = pathlib.Path(__file__).resolve().parent.parent
OUTDIR = REPO / "data" / "_qirqisani_build"

# Verified Cairo body-header start offsets (char index into the file). The body
# begins after the فهرس (table of contents). Each entry was confirmed to sit on
# an "الباب <ordinal> في <topic>" header line; internal "الباب" references
# between these offsets are correctly excluded by taking only this sequence.
CAIRO_BOUNDS = {
    1: 34880, 2: 38615, 3: 55245, 4: 98051, 5: 118945, 6: 120305,
    7: 121781, 8: 122647, 9: 130187, 10: 131350, 11: 138147, 12: 140109,
    13: 140512, 14: 144029, 15: 146689, 16: 148250, 17: 149365, 18: 150468,
    19: 152148, 20: 160804,  # 20 = start of Maqala II (end sentinel)
}

MARKER_RE = re.compile(r"\[\{\d+\}\]")           # English footnote markers [{1}]
APPARATUS_RE = re.compile(r"^\s*[\(（][\d٠-٩]+[\)）]")  # footnote lines "(١) ..."
RUNHEAD_RE = re.compile(r"الأنوار\s*والمراقب|الأنواروالمراقب")
PAGEMARK_RE = re.compile(r"^\s*[|\[]?\s*[طص]?\s*[\d٠-٩]+\s*[)\]]?\s*$")


def strip_markers(s):
    return MARKER_RE.sub("", s or "").strip()


def load_xlsx():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Text Content"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    ci = {name: i for i, name in enumerate(hdr)}
    idx_c, en_c = ci["Index"], ci["English"]
    babs = {}  # bab -> {"title_en":..., "segments":[{seg,en}]}
    for r in rows[1:]:
        raw_idx = r[idx_c]
        if not raw_idx:
            continue
        idx = str(raw_idx).strip()
        m_hdr = re.fullmatch(r"I\.(\d+)", idx)
        m_seg = re.fullmatch(r"I\.(\d+)\.(\d+)", idx)
        if m_hdr:
            b = int(m_hdr.group(1))
            babs.setdefault(b, {"title_en": "", "segments": []})
            babs[b]["title_en"] = strip_markers(r[en_c])
        elif m_seg:
            b, s = int(m_seg.group(1)), int(m_seg.group(2))
            babs.setdefault(b, {"title_en": "", "segments": []})
            babs[b]["segments"].append({"seg": s, "en": strip_markers(r[en_c])})
        # else: ignore non-conforming index rows
    for b in babs:
        babs[b]["segments"].sort(key=lambda x: x["seg"])
    return babs


def load_footnotes():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Footnotes"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        out.append({hdr[i]: r[i] for i in range(len(hdr))})
    return out


def clean_cairo_span(span):
    """Light pre-clean of a Cairo body span: drop editorial-apparatus lines,
    running heads, and bare page markers. Does NOT touch in-text OCR errors —
    that is the reconstruction agent's job."""
    keep = []
    for ln in span.split("\n"):
        t = ln.strip()
        if not t:
            continue
        if APPARATUS_RE.match(t):
            continue
        if RUNHEAD_RE.search(t) and len(t) < 40:
            continue
        if PAGEMARK_RE.match(t):
            continue
        keep.append(t)
    return "\n".join(keep)


def main():
    for p in (XLSX, CAIRO):
        if not p.exists():
            print(f"FATAL: missing {p}", file=sys.stderr)
            sys.exit(1)

    babs = load_xlsx()
    cairo = CAIRO.read_text(encoding="utf-8", errors="replace")

    # sanity: each mapped boundary should sit near an "الباب" header
    for b, off in CAIRO_BOUNDS.items():
        if b == 20:
            continue
        head = cairo[off:off + 30].replace("\n", " ")
        if "الباب" not in head and "الباب" not in cairo[off:off + 6]:
            print(f"WARN: bab {b} @ {off} does not start with الباب: {head!r}", file=sys.stderr)

    OUTDIR.mkdir(parents=True, exist_ok=True)
    skeleton = {}
    for b in sorted(babs):
        skeleton[b] = {"title_en": babs[b]["title_en"],
                       "segments": babs[b]["segments"]}

    (REPO / "data" / "_qirqisani_english_skeleton.json").write_text(
        json.dumps(skeleton, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    fns = load_footnotes()
    (REPO / "data" / "_qirqisani_footnotes.json").write_text(
        json.dumps(fns, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    built = []
    for b in range(2, 20):
        if b not in babs:
            print(f"WARN: bab {b} absent from xlsx", file=sys.stderr)
            continue
        span = cairo[CAIRO_BOUNDS[b]:CAIRO_BOUNDS[b + 1]]
        bundle = {
            "bab": b,
            "title_en": babs[b]["title_en"],
            "n_segments": len(babs[b]["segments"]),
            "segments": babs[b]["segments"],
            "cairo_arabic": clean_cairo_span(span),
        }
        (OUTDIR / f"bab{b}.json").write_text(
            json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        built.append((b, len(babs[b]["segments"])))

    print("Built bab bundles (bab: #segments):")
    for b, n in built:
        print(f"  bab {b:2d}: {n}")
    print(f"Total body segments (babs 2–19): {sum(n for _, n in built)}")
    print(f"Footnotes dumped: {len(fns)}")


if __name__ == "__main__":
    main()
